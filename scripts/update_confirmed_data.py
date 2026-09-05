"""Apply confirmed September 2026 business data to the source workbook."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CONFIRMED_FABRICS = {
    "Каспиан/Довер": {
        "width_cm": 280,
        "rrc": 1470,
        "comment": "Для моделей с тканью Каспиан использовать Каспиан/Довер.",
    },
    "Эйприл": {
        "width_cm": 300,
        "rrc": 390,
        "comment": "Ширина и розничная цена подтверждены вручную.",
    },
    "Сайфер": {
        "width_cm": 300,
        "rrc": 2680,
        "comment": "Ширина и розничная цена подтверждены вручную.",
    },
    "Луара": {
        "width_cm": 300,
        "rrc": 2800,
        "comment": "Ширина и розничная цена подтверждены вручную.",
    },
    "Сири": {
        "width_cm": 180,
        "rrc": 1510,
        "comment": "Розничная цена 1510 ₽/м подтверждена вручную.",
    },
    "Репаблик": {
        "width_cm": 300,
        "rrc": 1240,
        "comment": "Для расчёта подтверждена ширина 300 см.",
    },
}


def _headers(worksheet: Any) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }


def update_workbook(path: Path) -> None:
    workbook = load_workbook(path)

    fabrics_sheet = workbook["Ткани"]
    fabric_columns = _headers(fabrics_sheet)
    fabric_rows = {
        fabrics_sheet.cell(row=row, column=fabric_columns["Ткань"]).value: row
        for row in range(2, fabrics_sheet.max_row + 1)
    }
    for name, values in CONFIRMED_FABRICS.items():
        row = fabric_rows.get(name)
        if row is None:
            row = fabrics_sheet.max_row + 1
            fabrics_sheet.cell(row=row, column=fabric_columns["Ткань"], value=name)
            fabrics_sheet.cell(
                row=row,
                column=fabric_columns["Используется в MVP"],
                value="Да",
            )
            fabrics_sheet.cell(
                row=row,
                column=fabric_columns["Статус"],
                value="OK",
            )
        fabrics_sheet.cell(
            row=row,
            column=fabric_columns["Ширина, см"],
            value=values["width_cm"],
        )
        fabrics_sheet.cell(
            row=row,
            column=fabric_columns["Цена RRC, ₽/м"],
            value=values["rrc"],
        )
        fabrics_sheet.cell(
            row=row,
            column=fabric_columns["Используется в MVP"],
            value="Да",
        )
        fabrics_sheet.cell(
            row=row,
            column=fabric_columns["Комментарий"],
            value=values["comment"],
        )

    master_sheet = workbook["Мастер"]
    master_columns = _headers(master_sheet)
    for row in range(2, master_sheet.max_row + 1):
        fabric_cell = master_sheet.cell(
            row=row,
            column=master_columns["Ткань"],
        )
        if fabric_cell.value == "Каспиан":
            fabric_cell.value = "Каспиан/Довер"

        model = master_sheet.cell(
            row=row,
            column=master_columns["Модель"],
        ).value
        product_type = master_sheet.cell(
            row=row,
            column=master_columns["Тип изделия"],
        ).value
        if product_type == "Шторы" and model in {"Вилла", "Моррис"}:
            master_sheet.cell(
                row=row,
                column=master_columns["Статус доп. операции"],
                value="Проблемный",
            )
            master_sheet.cell(
                row=row,
                column=master_columns["Комментарий"],
                value=(
                    "Ткань временно исключена из расчёта. "
                    "Расчёт для этой модели пока недоступен."
                ),
            )

    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=Path("pasionaria_master_production_final_v2.xlsx"),
    )
    args = parser.parse_args()
    update_workbook(args.workbook)


if __name__ == "__main__":
    main()
