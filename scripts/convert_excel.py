"""Convert the PASIONARIA source workbook into deterministic JSON catalogs."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


REQUIRED_SHEETS = {
    "Мастер",
    "Ткани",
    "Операции",
    "Шаблоны",
    "Нюансы подтверждены",
}

PRODUCT_TYPES = {
    "Шторы": "curtain",
    "Римская штора": "roman",
}

RECIPE_OPERATION_KEYS = {
    "CURTAIN_BASE": [
        ("Шторы", "Крепление портьеры", "Шторная лента 6 см белая, без слоя"),
        ("Шторы", "Боковые швы", "2х2"),
        ("Шторы", "Нижний шов", "10х10"),
        ("Шторы", "Сборка шторной ленты", "Стандарт"),
        ("Шторы", "Упаковка портьеры", "Чехол ПВХ/спандбонд + картон"),
    ],
    "ROMAN_BASE": [
        ("Римская штора", "Пошив римской шторы", "Стандартная сборка"),
        ("Римская штора", "Механизм", "Стандарт"),
        ("Римская штора", "Крепления к механизму", "Кулиска + Кольца"),
        ("Римская штора", "Монтаж римской шторы", "Стандарт"),
    ],
}


def rows_as_dicts(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if not rows:
        return []
    headers = rows[0]
    return [
        {str(header): value for header, value in zip(headers, row)}
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def find_exact_operation(
    operations: Iterable[dict[str, Any]],
    key: tuple[str, str, str],
) -> dict[str, Any]:
    category, group, variant = key
    matches = [
        operation
        for operation in operations
        if (
            operation["category_source"] == category
            and operation["group"] == group
            and operation["variant"] == variant
        )
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one operation for {key!r}, found {len(matches)}")
    return matches[0]


def convert(workbook_path: Path, output_dir: Path) -> None:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    missing_sheets = REQUIRED_SHEETS.difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Workbook is missing sheets: {sorted(missing_sheets)}")

    source = workbook_path.name
    source_operations = rows_as_dicts(workbook, "Операции")
    operations: list[dict[str, Any]] = []
    for index, row in enumerate(source_operations, start=1):
        operations.append(
            {
                "id": f"OP_{index:03d}",
                "category_source": row["Категория"],
                "group": row["Группа операции"],
                "variant": row["Вариант"],
                "unit": row["Ед."],
                "tariff_rub": row["Тариф"],
                "calculation_basis": row["База расчета"],
                "formula": row["Формула для движка"],
                "source_page": row["Стр. PDF"],
                "comment": row["Комментарий"],
            }
        )

    source_master = rows_as_dicts(workbook, "Мастер")
    master_records: list[dict[str, Any]] = []
    for row in source_master:
        source_status = row["Статус доп. операции"]
        status = "problematic" if source_status == "Проблемный" else "active"
        extra_code = row["Доп. операция"]
        extra_operation_id = None
        if extra_code and source_status == "Подтверждено":
            group = (
                "Вышивка по всей площади"
                if str(extra_code).startswith("EMB_FULL_")
                else "Вышивка с одной стороны"
            )
            variant = (
                "Фито (мушки)"
                if extra_code == "EMB_FULL_ФИТО_МУШКИ"
                else row["Модель"]
            )
            extra_operation_id = find_exact_operation(
                operations,
                ("Шторы/Римские", group, variant),
            )["id"]

        master_records.append(
            {
                "product_type": PRODUCT_TYPES[row["Тип изделия"]],
                "product_type_source": row["Тип изделия"],
                "model": row["Модель"],
                "fabric_name": row["Ткань"],
                "recipe_code": row["Шаблон производства"],
                "additional_operation_code": extra_code,
                "additional_operation_id": extra_operation_id,
                "additional_operation_status_source": source_status,
                "status": status,
                "reason": row["Комментарий"],
            }
        )

    source_fabrics = rows_as_dicts(workbook, "Ткани")
    fabrics = [
        {
            "name": row["Ткань"],
            "base_article": row["Базовый артикул"],
            "width_cm": row["Ширина, см"],
            "base_price_rub_per_m": row["Цена, ₽/м"],
            "rrc_price_rub_per_m": row["Цена RRC, ₽/м"],
            "sku_count": row["Количество SKU"],
            "used_in_mvp": row["Используется в MVP"] == "Да",
            "status": row["Статус"],
            "comment": row["Комментарий"],
        }
        for row in source_fabrics
    ]

    source_templates = rows_as_dicts(workbook, "Шаблоны")
    templates_by_code = {row["Код шаблона"]: row for row in source_templates}
    recipes: list[dict[str, Any]] = []
    for code, operation_keys in RECIPE_OPERATION_KEYS.items():
        template = templates_by_code[code]
        recipes.append(
            {
                "code": code,
                "product_type": PRODUCT_TYPES[template["Тип изделия"]],
                "name": template["Название"],
                "operation_ids": [
                    find_exact_operation(operations, key)["id"]
                    for key in operation_keys
                ],
                "omitted_operations": (
                    [
                        {
                            "name": "Базовая обработка низа",
                            "reason": "Исключена по подтверждённому решению: тариф отсутствует.",
                        }
                    ]
                    if code == "ROMAN_BASE"
                    else []
                ),
                "source_operations": template["Базовые операции"],
                "source_rule": template["Правило расчета / тариф"],
            }
        )

    confirmed_nuances = rows_as_dicts(workbook, "Нюансы подтверждены")
    nuance_records = [
        {
            "product_type": PRODUCT_TYPES[row["Тип изделия"]],
            "model": row["Модель"],
            "fabric_name": row["Ткань"],
            "description": row["Подтвержденный нюанс"],
            "source_tariff": row["Тариф из PDF"],
            "status": row["Статус"],
        }
        for row in confirmed_nuances
    ]

    write_json(
        output_dir / "master.json",
        {"source": source, "records": master_records},
    )
    write_json(
        output_dir / "fabrics.json",
        {
            "source": source,
            "price_for_calculation": "rrc_price_rub_per_m",
            "records": fabrics,
        },
    )
    write_json(
        output_dir / "operations.json",
        {"source": source, "records": operations},
    )
    write_json(
        output_dir / "recipes.json",
        {
            "source": source,
            "recipes": recipes,
            "confirmed_nuances": nuance_records,
            "excluded_operation_groups": [
                {
                    "group": "Коэффициент типа ткани",
                    "reason": "Подтверждено: коэффициенты типа ткани не используются.",
                }
            ],
        },
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=Path("pasionaria_master_production_final_v2.xlsx"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("backend/app/data"),
    )
    args = parser.parse_args()
    convert(args.workbook, args.output_dir)


if __name__ == "__main__":
    main()
